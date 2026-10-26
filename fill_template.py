#!/usr/bin/env python3
"""Fill a copy of the deferred revenue template from one arr-check run.

    python3 fill_template.py out --out deferred_revenue_filled.xlsx

Reads deferred_rollforward.csv, deferred_accounts.csv, journal.csv and, when the run
wrote it, arr_to_revenue.csv from the run folder (and manifest.json for the as of
month, when it is there) and writes them into the yellow cells of a COPY of templates/deferred_revenue.xlsx. The template itself is
never changed, and every figure the template derives stays a formula: the closing, the
ledger difference, the journal tie and the tie out are computed by the spreadsheet,
not written here, so the filled file checks the run rather than repeating it.

The file is written without stored results; Excel and Google Sheets calculate on open.

REFUSES, writing nothing, when: a file is missing; the output already exists; the run
has more months, accounts or journal lines than the template holds; the as of month
is not a month of the roll forward; or a month's opening is not the month before's
closing (the template derives every opening after the first, so such a run could not
be represented faithfully).

Standard library plus openpyxl, which pilot.py already needs.
"""
import argparse, csv, json, os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "templates", "deferred_revenue.xlsx")
MONTHS, ACCOUNTS, LINES, STEPS = 60, 5000, 50000, 60


def refuse(msg):
    print(f"REFUSED, nothing written: {msg}")
    sys.exit(1)


def read(folder, name, required=True):
    path = os.path.join(folder, name)
    if not os.path.exists(path):
        if required:
            refuse(f"{name} is not in {folder}")
        return None
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def put(ws, ref, value):
    """Write a cell. Text is forced to be text: openpyxl stores any string that starts
    with "=" as a formula, so a step named "= supported by invoices", or a customer
    name or memo in an export that starts with "=", would otherwise become a formula
    in the filled copy."""
    ws[ref] = value
    if isinstance(value, str):
        ws[ref].data_type = "s"


def num(s):
    return None if s in ("", None) else float(s)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("run", help="an arr-check output folder")
    ap.add_argument("--out", required=True, help="the filled copy to write, .xlsx")
    ap.add_argument("--template", default=TEMPLATE)
    a = ap.parse_args()
    try:
        import openpyxl
    except ImportError:
        refuse("openpyxl is not installed for this Python (%s). Run: python3 -m pip install openpyxl" % sys.executable)

    if not a.out.lower().endswith(".xlsx"):
        refuse("--out must end in .xlsx")
    if os.path.exists(a.out):
        refuse(f"{a.out} already exists")
    if not os.path.exists(a.template):
        refuse(f"no template at {a.template}")
    roll = read(a.run, "deferred_rollforward.csv")
    accts = read(a.run, "deferred_accounts.csv")
    journal = read(a.run, "journal.csv")
    # arr_to_revenue.csv is newer than the other three; a run without it leaves the sheet blank.
    bridge = read(a.run, "arr_to_revenue.csv", required=False) or []
    for what, rows, cap in (("months", roll, MONTHS), ("accounts", accts, ACCOUNTS), ("journal lines", journal, LINES), ("bridge rows", bridge, STEPS)):
        if len(rows) > cap:
            refuse(f"{len(rows)} {what}; the template holds {cap}")
    if not roll:
        refuse("deferred_rollforward.csv has no months")
    for prev, cur in zip(roll, roll[1:]):
        if abs(num(cur["opening"]) - num(prev["closing"])) > 0.005:
            refuse(f"{cur['month']} opens at {cur['opening']}, {prev['month']} closed at {prev['closing']}")

    asof = roll[-1]["month"]
    man_path = os.path.join(a.run, "manifest.json")
    man = json.load(open(man_path)) if os.path.exists(man_path) else {}
    asof = man.get("asof_month") or asof
    if asof not in {r["month"] for r in roll}:
        refuse(f"the as of month {asof} is not a month of deferred_rollforward.csv")

    wb = openpyxl.load_workbook(a.template)
    ws = wb["Roll forward"]
    gl = "gl_balance" in roll[0]
    for i, r in enumerate(roll, start=2):
        put(ws, f"A{i}", r["month"])
        if i == 2:
            put(ws, "B2", num(r["opening"]))
        put(ws, f"C{i}", num(r["billings"]))
        put(ws, f"D{i}", num(r["recognised"]))
        put(ws, f"F{i}", num(r["closing_deferred"]))
        put(ws, f"G{i}", num(r["closing_unbilled"]))
        if gl:
            put(ws, f"I{i}", num(r["gl_balance"]))
    ws = wb["Accounts"]
    for i, r in enumerate(accts, start=2):
        put(ws, f"A{i}", r["key"])
        put(ws, f"B{i}", r["customer_name"])
        put(ws, f"C{i}", num(r["closing_deferred"]))
        put(ws, f"D{i}", num(r["closing_unbilled"]))
        put(ws, f"E{i}", num(r["months_of_deferred_cover"]))
    ws = wb["Journal"]
    for i, r in enumerate(journal, start=2):
        put(ws, f"A{i}", r["month"])
        put(ws, f"B{i}", r["customer_name"])
        put(ws, f"C{i}", r["debit_account"])
        put(ws, f"D{i}", r["credit_account"])
        put(ws, f"E{i}", num(r["amount"]))
        put(ws, f"F{i}", r["memo"])
    ws = wb["ARR to revenue"]
    for i, r in enumerate(bridge, start=2):
        kind = ("total" if r["step"].startswith("=") else "memo" if r["section"] == "memo"
                else "open" if not r["bucket"] else "step")
        put(ws, f"A{i}", int(r["order"]))
        put(ws, f"B{i}", r["section"])
        put(ws, f"C{i}", r["step"])
        put(ws, f"D{i}", r["bucket"])
        put(ws, f"E{i}", kind)
        # A total row's figure is the template's formula; the kit's figure goes beside it.
        put(ws, f"{'I' if kind == 'total' else 'F'}{i}", num(r["amount"]))
        put(ws, f"K{i}", r["accounts"])
    ws = wb["Tie out"]
    put(ws, "B3", asof)
    put(ws, "B26", " ".join(x for x in (os.path.basename(os.path.abspath(a.run)),
                                        f"kit {man['kit_version']}" if man.get("kit_version") else "",
                                        f"generated {man['generated_utc']} UTC" if man.get("generated_utc") else "") if x))
    wb.calculation.fullCalcOnLoad = True
    wb.save(a.out)
    print(f"wrote {a.out}: as of {asof}; {len(roll)} months, {len(accts)} accounts, {len(journal)} journal lines, {len(bridge)} bridge rows"
          f"{'' if gl else '; no ledger balance in this run, so Ledger balance and Difference stay blank'}")


if __name__ == "__main__":
    main()
